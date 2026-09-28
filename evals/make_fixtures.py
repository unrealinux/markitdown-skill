#!/usr/bin/env python3
"""Generate the fixtures used by evals/run_evals.py.

Standard library only, so the evals are reproducible on any machine:

    py evals/make_fixtures.py      # or python3 evals/make_fixtures.py

Rebuilds evals/fixtures/ from scratch every run.
"""

import os
import shutil
import struct
import subprocess
import sys
import wave
import zipfile
import zlib

HERE = os.path.dirname(os.path.abspath(__file__))
FIXTURES = os.path.join(HERE, "fixtures")

# xlsx and pptx fixtures are built with the real Office libraries, because a
# hand-rolled OOXML package has to satisfy every consumer (openpyxl, python-pptx)
# far more strictly than the zipfile-based .docx fixture does.
OFFICE_SCRIPT = '''
import os, sys
out = sys.argv[1]

from openpyxl import Workbook
book = Workbook()
sheet = book.active
sheet.title = "Sales"
sheet.append(["Region", "Sales"])
sheet.append(["North", "120"])
sheet.append(["South", "90"])

# A second sheet: the converter reads sheet_name=None, so every sheet lands in
# the markdown under its own ## heading. Asserted by eval 8.
stock = book.create_sheet("Inventory")
stock.append(["Item", "Qty"])
stock.append(["Widget", "7"])
book.save(os.path.join(out, "sample.xlsx"))

from pptx import Presentation
from pptx.util import Inches
prs = Presentation()
first = prs.slides.add_slide(prs.slide_layouts[5])
first.shapes.title.text = "PPTX Fixture Title"
second = prs.slides.add_slide(prs.slide_layouts[1])
second.shapes.title.text = "Second Slide"
second.placeholders[1].text = "Second slide body with office fixture"
prs.save(os.path.join(out, "sample.pptx"))
print("office fixtures written")
'''


def find_uv():
    """Same discovery order as scripts/convert.py."""
    for cmd in [
        os.path.expanduser("~/.local/bin/uv"),
        "uv",
        "/usr/local/bin/uv",
        "/opt/homebrew/bin/uv",
        os.path.join(os.environ.get("APPDATA", ""), "Python", "Python314", "Scripts", "uv.exe"),
        os.path.expanduser("~/AppData/Roaming/Python/Python314/Scripts/uv.exe"),
        os.path.expanduser("~/.local/bin/uv.exe"),
    ]:
        if cmd and (shutil.which(cmd) or os.path.exists(cmd)):
            return cmd
    return None


def write_office_fixtures():
    """Build sample.xlsx / sample.pptx through uv; returns True on success."""
    uv = find_uv()
    if not uv:
        print("uv not found: skipping sample.xlsx and sample.pptx")
        return False
    try:
        found = subprocess.run([uv, "python", "find", "3.12"],
                               capture_output=True, text=True, timeout=600)
    except (OSError, subprocess.SubprocessError):
        found = None
    python_path = None
    if found is not None and found.returncode == 0:
        for line in found.stdout.splitlines():
            if line.strip() and os.path.exists(line.strip()):
                python_path = line.strip()
                break

    cmd = [uv, "run", "--with", "openpyxl", "--with", "python-pptx"]
    if python_path:
        cmd += ["--python", python_path]
    cmd += ["python", "-c", OFFICE_SCRIPT, FIXTURES]

    env = os.environ.copy()
    env["PYTHONUTF8"] = "1"
    env["PYTHONIOENCODING"] = "utf-8"
    result = subprocess.run(cmd, capture_output=True, encoding="utf-8", errors="replace", env=env)
    if result.returncode != 0:
        print("could not build xlsx/pptx fixtures (network needed for uv extras):")
        for line in [l for l in (result.stderr or "").splitlines() if l.strip()][-3:]:
            print(f"  {line}")
        return False
    print((result.stdout or "").strip())
    return True


def setup_stdio():
    """Force UTF-8 so fixture content with emoji cannot break this script."""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="backslashreplace")
        except (AttributeError, ValueError):
            pass


def write_png(path, width=64, height=32, rgb=(0, 0, 0)):
    """Write a tiny solid-colour PNG (no text: markitdown has no OCR)."""
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(tag, data):
        crc = zlib.crc32(tag + data) & 0xFFFFFFFF
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", crc)

    blob = b"\x89PNG\r\n\x1a\n"
    blob += chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
    blob += chunk(b"IDAT", zlib.compress(raw, 9))
    blob += chunk(b"IEND", b"")
    with open(path, "wb") as handle:
        handle.write(blob)


CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""

ROOT_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""


def paragraph(text, style=None):
    properties = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return (
        f"<w:p>{properties}<w:r><w:t xml:space=\"preserve\">{text}</w:t></w:r></w:p>"
    )


def table(rows):
    cells = "".join(
        "<w:tr>"
        + "".join(
            f'<w:tc><w:p><w:r><w:t xml:space="preserve">{value}</w:t></w:r></w:p></w:tc>'
            for value in row
        )
        + "</w:tr>"
        for row in rows
    )
    return f"<w:tbl>{cells}</w:tbl>"


def write_docx(path):
    """Build a minimal but valid .docx (heading, paragraph, table)."""
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body>"
        + paragraph("Fixture Heading", style="Heading1")
        + paragraph("Fixture paragraph with 中文 and 😀 preserved.")
        + table([["Alpha", "Beta"], ["1", "2"]])
        + "</w:body></w:document>"
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", CONTENT_TYPES)
        archive.writestr("_rels/.rels", ROOT_RELS)
        archive.writestr("word/document.xml", document)


FOOTNOTE_CONTENT_TYPES = CONTENT_TYPES.replace(
    "</Types>",
    '<Override PartName="/word/footnotes.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.'
    'wordprocessingml.footnotes+xml"/>\n</Types>',
)

FOOTNOTE_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>
</Relationships>"""

FOOTNOTE_DOCUMENT = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    "<w:body>"
    '<w:p><w:r><w:t xml:space="preserve">Body text with a note</w:t></w:r>'
    '<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr>'
    '<w:footnoteReference w:id="1"/></w:r></w:p>'
    "</w:body></w:document>"
)

FOOTNOTE_PART = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
<w:footnote w:id="1"><w:p><w:r><w:t xml:space="preserve">FOOTNOTE-MARKER-TEXT</w:t></w:r></w:p></w:footnote>
</w:footnotes>"""


def write_footnote_docx(path):
    """A .docx holding one real footnote.

    markitdown renders it as an inline `[[1]](#footnote-1)` reference plus a
    trailing ordered list, NOT as `[^1]`; eval 19 pins that down.
    """
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", FOOTNOTE_CONTENT_TYPES)
        archive.writestr("_rels/.rels", ROOT_RELS)
        archive.writestr("word/_rels/document.xml.rels", FOOTNOTE_RELS)
        archive.writestr("word/document.xml", FOOTNOTE_DOCUMENT)
        archive.writestr("word/footnotes.xml", FOOTNOTE_PART)


def write_pdf(path, text="PDF fixture text 12345"):
    """Write a minimal single-page PDF with a real text layer (no OCR needed)."""
    stream = f"BT /F1 14 Tf 20 60 Td ({text}) Tj ET".encode("latin-1")
    bodies = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 120] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
        + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(bodies, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n"

    startxref = len(out)
    out += f"xref\n0 {len(bodies) + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {len(bodies) + 1} /Root 1 0 R >>\n"
        f"startxref\n{startxref}\n%%EOF\n"
    ).encode("ascii")

    with open(path, "wb") as handle:
        handle.write(bytes(out))


def write_zip(path, name="inner.txt", text="zip fixture content\n"):
    """One-entry archive. markitdown's ZipConverter unpacks it itself."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(name, text)


def write_wav(path, seconds=0.1, rate=8000):
    """A silent WAV.

    markitdown returns empty text for it: no exiftool for metadata, and the
    audio-transcription extra (speech_recognition) is never installed by the
    wrapper. That is the behaviour eval 17 pins down.
    """
    with wave.open(path, "wb") as handle:
        handle.setnchannels(1)
        handle.setsampwidth(1)
        handle.setframerate(rate)
        handle.writeframes(b"\x80" * int(seconds * rate))


def main():
    setup_stdio()
    if os.path.isdir(FIXTURES):
        shutil.rmtree(FIXTURES)
    os.makedirs(FIXTURES)

    with open(os.path.join(FIXTURES, "sample.txt"), "w", encoding="utf-8") as handle:
        handle.write("中文测试 plain text fixture\nemoji 😀 arrow → circled ①\n")

    with open(os.path.join(FIXTURES, "sample.csv"), "w", encoding="utf-8") as handle:
        handle.write("name,score\nAlpha,1\nBeta,2\n")

    with open(os.path.join(FIXTURES, "sample.html"), "w", encoding="utf-8") as handle:
        handle.write(
            "<html><head><title>Fixture</title>"
            "<script>console.log('should be stripped');</script>"
            "<style>body{color:red}</style></head>"
            "<body><h1>Fixture Heading</h1>"
            "<p>HTML paragraph with 中文.</p>"
            "<table><tr><th>Alpha</th><th>Beta</th></tr>"
            "<tr><td>1</td><td>2</td></tr></table></body></html>"
        )

    write_docx(os.path.join(FIXTURES, "sample.docx"))
    write_footnote_docx(os.path.join(FIXTURES, "sample_footnote.docx"))
    write_pdf(os.path.join(FIXTURES, "sample.pdf"))
    write_png(os.path.join(FIXTURES, "sample.png"))
    write_zip(os.path.join(FIXTURES, "sample.zip"))
    write_wav(os.path.join(FIXTURES, "sample.wav"))
    write_office_fixtures()

    tree = os.path.join(FIXTURES, "tree")
    for relative, text in [
        ("top.txt", "top level file\n"),
        ("a/report.txt", "A report\n"),
        ("b/report.txt", "B report\n"),
        ("b/sub/deep.txt", "deep note\n"),
        (".hidden/skipme.txt", "hidden file should be skipped\n"),
        ("notes.md", "markdown is not converted in directory mode\n"),
    ]:
        target = os.path.join(tree, relative)
        os.makedirs(os.path.dirname(target), exist_ok=True)
        with open(target, "w", encoding="utf-8") as handle:
            handle.write(text)

    # A legacy .doc: markitdown 0.1.8 has no converter for it and no LibreOffice
    # path, so directory mode must report it as skipped instead of failing.
    with open(os.path.join(tree, "legacy.doc"), "wb") as handle:
        handle.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(range(256)) * 4)

    # Failure path: one readable file next to a .docx whose XML is garbage.
    broken = os.path.join(FIXTURES, "broken")
    os.makedirs(broken, exist_ok=True)
    with open(os.path.join(broken, "good.txt"), "w", encoding="utf-8") as handle:
        handle.write("good content survives a sibling failure\n")
    with zipfile.ZipFile(os.path.join(broken, "bad.docx"), "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        )
        archive.writestr("word/document.xml", "<<<not xml at all")

    # Two sources, one stem: report.txt and report.csv both used to write report.md.
    collide = os.path.join(FIXTURES, "collide")
    os.makedirs(collide, exist_ok=True)
    with open(os.path.join(collide, "report.txt"), "w", encoding="utf-8") as handle:
        handle.write("from the text file\n")
    with open(os.path.join(collide, "report.csv"), "w", encoding="utf-8") as handle:
        handle.write("source,alpha\n")
    with open(os.path.join(collide, "other.txt"), "w", encoding="utf-8") as handle:
        handle.write("unrelated\n")

    # Empty-output path: an image with no text layer next to a real text file.
    # Batch mode must warn (⚠️ 0 chars) instead of claiming a clean conversion.
    scan = os.path.join(FIXTURES, "scan")
    os.makedirs(scan, exist_ok=True)
    write_png(os.path.join(scan, "scan.png"))
    with open(os.path.join(scan, "note.txt"), "w", encoding="utf-8") as handle:
        handle.write("scan folder note\n")

    print(f"fixtures written to {FIXTURES}")


if __name__ == "__main__":
    main()
